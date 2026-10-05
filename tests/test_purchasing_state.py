"""
Tests for workflow_runner.py's purchasing-handoff functions:
derive_purchasing_status, build_purchasing_summary, regenerate_purchasing,
and the per-artifact freshness/state model (artifact_status) that backs
them.

Both DigiKey and Mouser purchasing artifacts are now plain, local,
side-effect-free operations -- a CSV file write (digikey_list_csv/
mouser_export_csv) or a pure deterministic URL derivation
(digikey_cart_link, via src/digikey_fastadd.py). Neither makes a vendor
API call, so neither needs an enable flag or mocking to run safely here.

DigiKey MyList creation (src/digikey_mylists_client.py) still exists as a
module, but workflow_runner.py no longer imports or calls it anywhere in
the normal run/regenerate path -- see
Documentation/AutoBOM_Direct_Cart_Handoff_Change_Risk_Audit.md. Tests that
need to assert "no vendor call happened" patch
digikey_mylists_client.create_digikey_mylist_from_bom directly at its
defining module (NOT workflow_runner.create_digikey_mylist_from_bom --
that name doesn't exist in workflow_runner's namespace anymore and
patching it there raises AttributeError).
"""

import os
import shutil
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import pandas as pd

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SRC_FOLDER = PROJECT_ROOT / "src"
sys.path.insert(0, str(SRC_FOLDER))

from purchasing_lines import PurchasingLineResult
from review_actions import apply_replacement
from run_workspace import build_run_workspace, save_uploaded_file
from sourcing_engine import SupplierResult
from workflow_runner import (
    build_purchasing_summary,
    derive_purchasing_status,
    refresh_response_snapshot,
    regenerate_purchasing,
    run_bom_workflow,
)


CLEAN_BOM_CSV = (
    "Designator,Quantity,Manufacturer,Part Number,Description\n"
    "R1 R2 R3,3,Yageo,RC0603FR-0710KL,10k Resistor 1% 0603\n"
    "C1 C2,2,Murata,GRM188R71C105KA12D,1uF Capacitor 16V 0603\n"
)

REVIEW_BOM_CSV = (
    "Designator,Quantity,Manufacturer,Part Number,Description\n"
    "U1,1,STMicroelectronics,,Missing part number\n"
)


class PurchasingStateTestCase(unittest.TestCase):
    def setUp(self):
        self.temp_root = Path(tempfile.mkdtemp(prefix="autobom_purchasing_test_"))
        self.addCleanup(shutil.rmtree, self.temp_root, ignore_errors=True)

        self._saved_env = {}
        for key in ("SUPPLIER_LOOKUP_CACHE_ENABLED",):
            self._saved_env[key] = os.environ.get(key)
        os.environ["SUPPLIER_LOOKUP_CACHE_ENABLED"] = "false"
        self.addCleanup(self._restore_env)

    def _restore_env(self):
        for key, value in self._saved_env.items():
            if value is None:
                os.environ.pop(key, None)
            else:
                os.environ[key] = value

    def run_workflow(self, csv_text=CLEAN_BOM_CSV, project_name="PurchasingTest"):
        workspace = build_run_workspace(runs_root=self.temp_root)
        source_path = save_uploaded_file(workspace["input_dir"], "bom.csv", csv_text.encode("utf-8"))
        return run_bom_workflow(
            source_file_path=source_path,
            build_quantity=5,
            project_name=project_name,
            run_partsbox=False,
            run_sourcing=False,
            run_id=workspace["run_id"],
            output_folder=str(workspace["output_dir"]),
        )

    def make_mouser_sourced(self, active_bom_state, row_id, mpn="603-ABC123", qty="15"):
        active_bom_state.update_row(
            row_id,
            {
                "selected_supplier": "Mouser",
                "supplier_part_number": mpn,
                "sourcing_status": "sourced_mouser",
                "supplier_order_qty": qty,
            },
        )

    def make_digikey_sourced(self, active_bom_state, row_id, mpn="296-XYZ", qty="10"):
        active_bom_state.update_row(
            row_id,
            {
                "selected_supplier": "DigiKey",
                "supplier_part_number": mpn,
                "sourcing_status": "sourced_digikey",
                "supplier_order_qty": qty,
            },
        )

    def make_digikey_sourced_missing_part_number(self, active_bom_state, row_id, qty="10"):
        """
        A row that IS sourced from DigiKey (selected_supplier/sourcing_status
        both set) but has no usable DigiKey supplier part number -- a
        genuine data problem, not something purchasing_lines.py should
        silently drop. Must surface as a PurchasingLineError (see H2 in
        the change-risk audit), never let the handoff report Ready.
        """
        active_bom_state.update_row(
            row_id,
            {
                "selected_supplier": "DigiKey",
                "supplier_part_number": "",
                "sourcing_status": "sourced_digikey",
                "supplier_order_qty": qty,
            },
        )

    def make_row_dnp(self, active_bom_state, row_id):
        """
        An explicit operator "don't populate this line" decision --
        purchasing_lines.build_purchasing_lines() must silently exclude
        it (no PurchasingLineError), whatever its supplier/sourcing
        fields say.
        """
        active_bom_state.update_row(row_id, {"dnp": True})


class _FakeSupplierClient:
    def __init__(self, result):
        self.result = result

    def find_best_match_for_row(self, row):
        return self.result


class DerivePurchasingStatusTests(PurchasingStateTestCase):
    def test_not_available_without_active_state(self):
        self.assertEqual(derive_purchasing_status({"_active_bom_state": None}), "Not Available")

    def test_needs_review_when_rows_pending(self):
        result = self.run_workflow(csv_text=REVIEW_BOM_CSV)
        self.assertEqual(derive_purchasing_status(result), "Needs Review")

    def test_not_yet_generated_when_clean_and_untouched(self):
        result = self.run_workflow(csv_text=CLEAN_BOM_CSV)
        self.assertEqual(derive_purchasing_status(result), "Not Yet Generated")

    def test_stale_after_mutation_takes_priority(self):
        result = self.run_workflow(csv_text=CLEAN_BOM_CSV)
        result["artifact_status"]["cleaned_csv"] = "current"
        refresh_response_snapshot(result)
        self.assertEqual(derive_purchasing_status(result), "Stale — BOM Changed")

    def test_zero_required_items_after_regeneration_is_ready_for_handoff(self):
        # Nothing ever attempted (fresh run) reports "Not Yet Generated";
        # once an explicit regeneration has actually run and neither
        # supplier has any purchasing lines required, there is nothing
        # left to do -- that must read as ready, not stuck at "not yet".
        result = self.run_workflow(csv_text=CLEAN_BOM_CSV)

        regenerate_purchasing(result)

        self.assertEqual(derive_purchasing_status(result), "Ready for Handoff")

    def test_mouser_sourced_items_never_block_readiness(self):
        """
        DigiKey is the only supplier with zero purchasing lines here, so
        neither of its REQUIRED_ARTIFACT_KEYS entries is required --
        having Mouser-sourced items and zero DigiKey-required items must
        still read as Ready for Handoff.
        """
        result = self.run_workflow(csv_text=CLEAN_BOM_CSV)
        active_bom_state = result["_active_bom_state"]
        row_id = active_bom_state.rows()[0]["row_id"]
        self.make_mouser_sourced(active_bom_state, row_id)

        regenerate_purchasing(result)

        self.assertEqual(result["artifact_status"]["mouser_export_csv"], "current")
        self.assertEqual(result["artifact_status"]["digikey_list_csv"], "not_required")
        self.assertEqual(result["artifact_status"]["digikey_cart_link"], "not_required")
        self.assertEqual(derive_purchasing_status(result), "Ready for Handoff")

    def test_mouser_export_failure_does_not_block_ready_for_handoff_status(self):
        """
        A failed Mouser export disables that one download (see the
        artifact_status assertion) but must not itself block "Ready for
        Handoff" here -- there are zero Mouser-sourced lines in this BOM,
        so mouser_export_csv is not required at all; readiness is judged
        only against REQUIRED_ARTIFACT_KEYS entries that are actually
        required for the current BOM contents.
        """
        result = self.run_workflow(csv_text=CLEAN_BOM_CSV)

        with patch("workflow_runner.export_mouser_list") as mock_export:
            mock_export.side_effect = RuntimeError("disk full")
            regenerate_purchasing(result)

        self.assertEqual(result["artifact_status"]["mouser_export_csv"], "failed")
        self.assertEqual(derive_purchasing_status(result), "Ready for Handoff")

    def test_mouser_and_digikey_both_sourced_and_valid_is_ready_for_handoff(self):
        result = self.run_workflow(csv_text=CLEAN_BOM_CSV)
        active_bom_state = result["_active_bom_state"]
        rows = active_bom_state.rows()
        self.make_mouser_sourced(active_bom_state, rows[0]["row_id"])
        self.make_digikey_sourced(active_bom_state, rows[1]["row_id"])

        regenerate_purchasing(result)

        self.assertEqual(result["artifact_status"]["mouser_export_csv"], "current")
        self.assertEqual(result["artifact_status"]["digikey_list_csv"], "current")
        self.assertEqual(result["artifact_status"]["digikey_cart_link"], "current")
        self.assertEqual(derive_purchasing_status(result), "Ready for Handoff")


class BuildPurchasingSummaryTests(PurchasingStateTestCase):
    def test_summary_counts_sourced_rows_only(self):
        result = self.run_workflow(csv_text=CLEAN_BOM_CSV)
        active_bom_state = result["_active_bom_state"]
        row_id = active_bom_state.rows()[0]["row_id"]
        self.make_mouser_sourced(active_bom_state, row_id)

        summary = build_purchasing_summary(result)

        self.assertEqual(summary["mouser"]["item_count"], 1)
        self.assertEqual(summary["mouser"]["total_quantity"], 15)
        self.assertEqual(summary["digikey"]["item_count"], 0)

    def test_unresolved_rows_excluded_from_counts(self):
        result = self.run_workflow(csv_text=REVIEW_BOM_CSV)
        summary = build_purchasing_summary(result)

        self.assertEqual(summary["mouser"]["item_count"], 0)
        self.assertEqual(summary["digikey"]["item_count"], 0)
        self.assertEqual(summary["needs_review_count"], 1)

    def test_summary_reflects_replacement_resolution(self):
        result = self.run_workflow(csv_text=REVIEW_BOM_CSV)
        active_bom_state = result["_active_bom_state"]
        row_id = active_bom_state.rows()[0]["row_id"]

        mouser = _FakeSupplierClient(
            SupplierResult(
                supplier="Mouser",
                manufacturer="STMicroelectronics",
                mpn="FOUND123",
                stock=500,
                supplier_part_number="511-FOUND123",
            )
        )
        digikey = _FakeSupplierClient(None)
        apply_replacement(
            active_bom_state,
            row_id,
            {"source": "Mouser", "mpn": "FOUND123", "manufacturer": "STMicroelectronics"},
            mouser_client=mouser,
            digikey_client=digikey,
        )

        summary = build_purchasing_summary(result)
        self.assertEqual(summary["mouser"]["item_count"], 1)
        self.assertEqual(summary["needs_review_count"], 0)

    def test_purchasing_item_counts_derive_from_current_canonical_bom(self):
        """
        build_purchasing_summary() must read item counts from the LIVE
        ActiveBomState dataframe on every call, not from a snapshot taken
        at some earlier point -- a second mutation must be reflected
        immediately, with no regeneration required just to see updated
        counts (regeneration is only needed to actually create the
        cart/list artifacts).
        """
        result = self.run_workflow(csv_text=CLEAN_BOM_CSV)
        active_bom_state = result["_active_bom_state"]
        rows = active_bom_state.rows()

        self.assertEqual(build_purchasing_summary(result)["mouser"]["item_count"], 0)

        self.make_mouser_sourced(active_bom_state, rows[0]["row_id"])
        self.assertEqual(build_purchasing_summary(result)["mouser"]["item_count"], 1)

        self.make_mouser_sourced(active_bom_state, rows[1]["row_id"], mpn="603-DEF456")
        self.assertEqual(build_purchasing_summary(result)["mouser"]["item_count"], 2)


class RegeneratePurchasingTests(PurchasingStateTestCase):
    def test_noop_without_active_state(self):
        response = {"_active_bom_state": None}
        self.assertIs(regenerate_purchasing(response), response)

    # --- No DigiKey MyLists API call anywhere in the normal flow ---

    def test_regenerate_purchasing_never_calls_digikey_mylists_even_with_digikey_rows(self):
        """
        DigiKey purchasing artifacts (CSV + FastAdd cart link) are a
        local file write and a pure URL derivation -- regenerate_purchasing()
        must never reach digikey_mylists_client.create_digikey_mylist_from_bom(),
        even when DigiKey-sourced rows exist. Patched at its defining
        module: workflow_runner.py no longer imports this name at all, so
        patching "workflow_runner.create_digikey_mylist_from_bom" would
        raise AttributeError.
        """
        result = self.run_workflow(csv_text=CLEAN_BOM_CSV)
        active_bom_state = result["_active_bom_state"]
        row_id = active_bom_state.rows()[0]["row_id"]
        self.make_digikey_sourced(active_bom_state, row_id)

        with patch("digikey_mylists_client.create_digikey_mylist_from_bom") as mock_mylist:
            regenerate_purchasing(result)

        mock_mylist.assert_not_called()

    def test_run_bom_workflow_never_calls_digikey_mylists(self):
        with patch("digikey_mylists_client.create_digikey_mylist_from_bom") as mock_mylist:
            self.run_workflow(csv_text=CLEAN_BOM_CSV)

        mock_mylist.assert_not_called()

    def test_plain_rerun_creates_no_vendor_or_file_side_effect(self):
        """
        Reading purchasing state (as every Streamlit rerun does, simply by
        re-executing the page function) must never re-export any artifact
        or call any vendor function -- only an explicit regenerate_purchasing()
        call may.
        """
        result = self.run_workflow(csv_text=CLEAN_BOM_CSV)
        active_bom_state = result["_active_bom_state"]
        row_id = active_bom_state.rows()[0]["row_id"]
        self.make_mouser_sourced(active_bom_state, row_id)

        with patch("workflow_runner.export_mouser_list") as mock_export_mouser, patch(
            "digikey_mylists_client.create_digikey_mylist_from_bom"
        ) as mock_mylist:
            for _ in range(3):
                build_purchasing_summary(result)
                derive_purchasing_status(result)

        mock_export_mouser.assert_not_called()
        mock_mylist.assert_not_called()

    # --- DigiKey FastAdd cart-link artifact (src/digikey_fastadd.py) ---

    def test_valid_digikey_row_produces_current_csv_and_cart_link(self):
        result = self.run_workflow(csv_text=CLEAN_BOM_CSV)
        active_bom_state = result["_active_bom_state"]
        row_id = active_bom_state.rows()[0]["row_id"]
        self.make_digikey_sourced(active_bom_state, row_id)

        regenerate_purchasing(result)

        self.assertEqual(result["artifact_status"]["digikey_list_csv"], "current")
        self.assertEqual(result["artifact_status"]["digikey_cart_link"], "current")
        summary = build_purchasing_summary(result)
        self.assertTrue(
            summary["digikey"]["cart_link_url"].startswith(
                "https://www.digikey.com/classic/ordering/fastadd.aspx"
            )
        )

    def test_digikey_row_missing_supplier_part_number_fails_and_blocks_ready(self):
        result = self.run_workflow(csv_text=CLEAN_BOM_CSV)
        active_bom_state = result["_active_bom_state"]
        row_id = active_bom_state.rows()[0]["row_id"]
        self.make_digikey_sourced_missing_part_number(active_bom_state, row_id)

        regenerate_purchasing(result)

        self.assertEqual(result["artifact_status"]["digikey_list_csv"], "failed")
        self.assertEqual(result["artifact_status"]["digikey_cart_link"], "failed")
        summary = build_purchasing_summary(result)
        self.assertTrue(summary["digikey"]["errors"])
        self.assertNotEqual(derive_purchasing_status(result), "Ready for Handoff")

    def test_dnp_digikey_row_is_silently_excluded(self):
        result = self.run_workflow(csv_text=CLEAN_BOM_CSV)
        active_bom_state = result["_active_bom_state"]
        row_id = active_bom_state.rows()[0]["row_id"]
        self.make_digikey_sourced(active_bom_state, row_id)
        self.make_row_dnp(active_bom_state, row_id)

        regenerate_purchasing(result)

        summary = build_purchasing_summary(result)
        self.assertEqual(summary["digikey"]["item_count"], 0)
        self.assertEqual(summary["digikey"]["errors"], [])
        self.assertEqual(result["artifact_status"]["digikey_list_csv"], "not_required")
        self.assertEqual(result["artifact_status"]["digikey_cart_link"], "not_required")
        self.assertEqual(derive_purchasing_status(result), "Ready for Handoff")

    def test_oversized_digikey_cart_link_falls_back_to_csv_and_stays_ready(self):
        """
        _is_required_artifact_valid() explicitly accepts STATUS_CSV_FALLBACK
        for digikey_cart_link as "current and valid" -- a deliberate,
        documented terminal state (H1/H6 in the change-risk audit), not a
        degraded one. As long as digikey_list_csv is current, an oversized
        (but otherwise fully valid) DigiKey purchasing set is still Ready
        for Handoff, and the oversized/truncated URL itself must never be
        exposed.
        """
        result = self.run_workflow(csv_text=CLEAN_BOM_CSV)
        active_bom_state = result["_active_bom_state"]

        bulk_rows = pd.DataFrame(
            [
                {
                    "selected_supplier": "DigiKey",
                    "sourcing_status": "sourced_digikey",
                    "supplier_part_number": f"296-{index:05d}-ND-XX",
                    "supplier_order_qty": "1",
                    "customer_reference": f"CREF-{index:04d}",
                    "mpn": f"MPN{index:05d}",
                    "manufacturer": "Vendor",
                }
                for index in range(150)
            ]
        )
        active_bom_state.add_rows(bulk_rows)

        regenerate_purchasing(result)

        self.assertEqual(result["artifact_status"]["digikey_list_csv"], "current")
        self.assertEqual(result["artifact_status"]["digikey_cart_link"], "csv_fallback")
        summary = build_purchasing_summary(result)
        self.assertEqual(summary["digikey"]["cart_link_url"], "")
        self.assertEqual(derive_purchasing_status(result), "Ready for Handoff")

    def test_digikey_csv_export_failure_blocks_ready_when_digikey_required(self):
        """
        Distinct from test_mouser_export_failure_does_not_block_ready_for_handoff_status:
        there a failed Mouser export didn't matter because there were zero
        Mouser-sourced lines (not required at all). Here there IS a valid
        DigiKey-sourced line, so DigiKey CSV export failing is a required
        artifact failing, and must block Ready.
        """
        result = self.run_workflow(csv_text=CLEAN_BOM_CSV)
        active_bom_state = result["_active_bom_state"]
        row_id = active_bom_state.rows()[0]["row_id"]
        self.make_digikey_sourced(active_bom_state, row_id)

        with patch("workflow_runner.export_digikey_list") as mock_export:
            mock_export.side_effect = RuntimeError("boom")
            regenerate_purchasing(result)

        self.assertEqual(result["artifact_status"]["digikey_list_csv"], "failed")
        self.assertNotEqual(derive_purchasing_status(result), "Ready for Handoff")

    def test_cart_link_url_is_cleared_after_a_failed_regeneration(self):
        """
        A previous successful FastAdd link must never be left displayed
        next to a failed current status (H6 in the change-risk audit) --
        once regeneration fails, digikey_cart_link.url must go back to "".
        """
        result = self.run_workflow(csv_text=CLEAN_BOM_CSV)
        active_bom_state = result["_active_bom_state"]
        row_id = active_bom_state.rows()[0]["row_id"]
        self.make_digikey_sourced(active_bom_state, row_id)

        regenerate_purchasing(result)
        self.assertEqual(result["artifact_status"]["digikey_cart_link"], "current")
        first_url = result["digikey_cart_link"]["url"]
        self.assertTrue(first_url)

        active_bom_state.update_row(row_id, {"supplier_part_number": ""})
        regenerate_purchasing(result)

        self.assertNotEqual(result["artifact_status"]["digikey_cart_link"], "current")
        self.assertEqual(result["digikey_cart_link"]["url"], "")

    def test_mouser_export_regeneration_records_failure_without_raising(self):
        result = self.run_workflow(csv_text=CLEAN_BOM_CSV)

        with patch("workflow_runner.export_mouser_list") as mock_export:
            mock_export.side_effect = RuntimeError("disk full")
            regenerate_purchasing(result)

        self.assertIn(
            "Mouser CSV export failed", " ".join(result["purchasing_regeneration_errors"])
        )
        self.assertEqual(result["artifact_status"]["mouser_export_csv"], "failed")

    def test_regeneration_reexports_current_csv_outputs(self):
        result = self.run_workflow(csv_text=CLEAN_BOM_CSV)
        active_bom_state = result["_active_bom_state"]
        row_id = active_bom_state.rows()[0]["row_id"]
        active_bom_state.update_row(row_id, {"mpn": "CHANGED123"})

        regenerate_purchasing(result)

        csv_path = Path(result["outputs"]["cleaned_csv"])
        self.assertTrue(csv_path.exists())
        self.assertIn("CHANGED123", csv_path.read_text())

    def test_replacement_then_regeneration_produces_updated_purchasing_items(self):
        result = self.run_workflow(csv_text=REVIEW_BOM_CSV)
        active_bom_state = result["_active_bom_state"]
        row_id = active_bom_state.rows()[0]["row_id"]

        mouser = _FakeSupplierClient(
            SupplierResult(
                supplier="Mouser",
                manufacturer="STMicroelectronics",
                mpn="FOUND123",
                stock=500,
                supplier_part_number="511-FOUND123",
            )
        )
        apply_replacement(
            active_bom_state,
            row_id,
            {"source": "Mouser", "mpn": "FOUND123", "manufacturer": "STMicroelectronics"},
            mouser_client=mouser,
            digikey_client=_FakeSupplierClient(None),
        )

        regenerate_purchasing(result)

        summary = build_purchasing_summary(result)
        self.assertEqual(summary["mouser"]["item_count"], 1)
        self.assertEqual(summary["status"], "Ready for Handoff")

    # --- Per-artifact freshness / failure-isolation tests (Phase 4 fix pass) ---

    # The old test_retry_after_digikey_failure_does_not_recreate_it_once_current
    # has no replacement here: it protected against re-invoking a REAL
    # vendor side effect (creating a duplicate DigiKey MyList) after a
    # retry succeeded. DigiKey's CSV and FastAdd cart link are a plain
    # file write and a pure URL derivation with no vendor side effect to
    # protect against duplicating -- see
    # test_supplier_exports_and_digikey_cart_link_rerun_on_every_regeneration
    # below for the new (opposite) expectation that they rerun every time.

    def test_supplier_exports_and_digikey_cart_link_rerun_on_every_regeneration(self):
        """
        Every purchasing artifact here (Mouser CSV, DigiKey CSV, DigiKey
        FastAdd cart link) is a plain file write or a pure URL derivation,
        not a live vendor call -- unlike the old DigiKey MyList creation,
        there is no duplicate-side-effect risk to guard against, so all
        three simply re-run on every explicit regeneration, exactly like
        sourcing_report/cleaned_csv already do.
        """
        result = self.run_workflow(csv_text=CLEAN_BOM_CSV)
        active_bom_state = result["_active_bom_state"]
        row_id = active_bom_state.rows()[0]["row_id"]
        self.make_mouser_sourced(active_bom_state, row_id)

        with patch("workflow_runner.export_mouser_list") as mock_export:
            mock_export.return_value = ("/tmp/mouser_export.csv", PurchasingLineResult())
            regenerate_purchasing(result)
            regenerate_purchasing(result)

        self.assertEqual(mock_export.call_count, 2)

    def test_cleaned_csv_failure_is_caught_and_reported(self):
        result = self.run_workflow(csv_text=CLEAN_BOM_CSV)

        with patch("workflow_runner.export_clean_bom") as mock_export:
            mock_export.side_effect = RuntimeError("disk full")
            response = regenerate_purchasing(result)

        self.assertIs(response, result)
        self.assertEqual(result["artifact_status"]["cleaned_csv"], "failed")
        self.assertIn("Cleaned CSV regeneration failed", " ".join(result["purchasing_regeneration_errors"]))
        # Other independent artifacts still regenerated successfully.
        self.assertEqual(result["artifact_status"]["sourcing_report"], "current")

    def test_one_export_failure_keeps_only_that_artifact_stale_or_failed(self):
        result = self.run_workflow(csv_text=CLEAN_BOM_CSV)

        with patch("workflow_runner.export_digikey_list") as mock_export:
            mock_export.side_effect = RuntimeError("boom")
            regenerate_purchasing(result)

        self.assertEqual(result["artifact_status"]["digikey_list_csv"], "failed")
        self.assertEqual(result["artifact_status"]["cleaned_csv"], "current")
        self.assertEqual(result["artifact_status"]["sourcing_report"], "current")

    def test_stale_workbook_regenerates_alongside_other_artifacts(self):
        """
        QA/QC deep-dive report, Q9: the multi-sheet cleaned workbook used
        to have no regeneration path and stayed permanently disabled
        after the first post-upload edit. It now regenerates from the
        current canonical state just like cleaned_csv.
        """
        result = self.run_workflow(csv_text=CLEAN_BOM_CSV)
        self.assertEqual(result["artifact_status"]["cleaned_workbook"], "current")

        active_bom_state = result["_active_bom_state"]
        row_id = active_bom_state.rows()[0]["row_id"]
        active_bom_state.update_row(row_id, {"mpn": "CHANGED123"})
        refresh_response_snapshot(result)
        self.assertEqual(result["artifact_status"]["cleaned_workbook"], "stale")

        regenerate_purchasing(result)

        self.assertEqual(result["artifact_status"]["cleaned_csv"], "current")
        self.assertEqual(result["artifact_status"]["sourcing_report"], "current")
        self.assertEqual(result["artifact_status"]["cleaned_workbook"], "current")

    def test_regenerated_workbook_reflects_the_current_mutated_bom(self):
        import openpyxl

        result = self.run_workflow(csv_text=CLEAN_BOM_CSV)
        active_bom_state = result["_active_bom_state"]
        row_id = active_bom_state.rows()[0]["row_id"]
        active_bom_state.update_row(row_id, {"mpn": "CHANGED123"})
        refresh_response_snapshot(result)

        regenerate_purchasing(result)

        workbook_path = result["outputs"]["cleaned_workbook"]
        workbook = openpyxl.load_workbook(workbook_path)
        self.assertEqual(
            set(workbook.sheetnames), {"Clean_BOM", "Review_Items", "Column_Mapping", "Warnings"}
        )
        clean_bom_sheet = workbook["Clean_BOM"]
        mpn_values = [
            cell.value for row in clean_bom_sheet.iter_rows(min_row=2) for cell in row
        ]
        self.assertIn("CHANGED123", mpn_values)
        # Column_Mapping/Warnings describe the ORIGINAL upload's raw-to-
        # clean transformation, which cannot be recomputed post-mutation
        # -- present as empty (headers only), never fabricated.
        self.assertEqual(workbook["Column_Mapping"].max_row, 1)
        self.assertEqual(workbook["Warnings"].max_row, 1)

    def test_partsbox_unmatched_csv_regenerated_or_remains_unavailable(self):
        result = self.run_workflow(csv_text=CLEAN_BOM_CSV)
        result["artifact_status"]["partsbox_unmatched_csv"] = "current"

        active_bom_state = result["_active_bom_state"]
        row_id = active_bom_state.rows()[0]["row_id"]
        active_bom_state.update_row(row_id, {"mpn": "CHANGED123"})
        refresh_response_snapshot(result)
        self.assertEqual(result["artifact_status"]["partsbox_unmatched_csv"], "stale")

        regenerate_purchasing(result)

        # regenerate_purchasing() never re-derives PartsBox unmatched rows
        # (that requires a fresh, side-effecting project/get-entries call)
        # -- it must remain disabled, not silently re-enabled.
        self.assertEqual(result["artifact_status"]["partsbox_unmatched_csv"], "stale")

    def test_stale_downloads_never_become_enabled_merely_because_another_artifact_succeeded(self):
        result = self.run_workflow(csv_text=CLEAN_BOM_CSV)
        active_bom_state = result["_active_bom_state"]
        row_id = active_bom_state.rows()[0]["row_id"]
        active_bom_state.update_row(row_id, {"mpn": "CHANGED123"})
        refresh_response_snapshot(result)

        with patch("workflow_runner.export_sourcing_report") as mock_export:
            mock_export.side_effect = RuntimeError("boom")
            regenerate_purchasing(result)

        # cleaned_csv succeeded and is current; sourcing_report failed and
        # must stay disabled regardless of cleaned_csv's success.
        self.assertEqual(result["artifact_status"]["cleaned_csv"], "current")
        self.assertEqual(result["artifact_status"]["sourcing_report"], "failed")


class LegacyStateCompatibilityTests(PurchasingStateTestCase):
    """
    H9 in the change-risk audit: an existing Streamlit session that
    survived a hot reload (or a test/tool that injects an old response
    shape) may still carry pre-migration MyList-shaped keys --
    artifact_status["digikey_mylist"] and a top-level "purchasing" dict
    with DigiKey MyList creation/list_id/share_id fields. Neither
    derive_purchasing_status() nor build_purchasing_summary() may crash
    on that, and neither may let those legacy keys influence readiness or
    the summary shape -- readiness/summary must be judged purely from the
    current REQUIRED_ARTIFACT_KEYS/artifact_status model.
    """

    def test_legacy_mylist_shaped_response_does_not_crash_or_influence_new_readiness(self):
        result = self.run_workflow(csv_text=CLEAN_BOM_CSV)

        # Simulate a pre-migration in-memory response surviving a hot
        # reload: legacy keys added on top of an otherwise-fresh,
        # untouched-since-run response.
        result["artifact_status"]["digikey_mylist"] = "current"
        result["purchasing"] = {
            "mouser": None,
            "digikey": {"created": True, "list_id": "OLD-LIST-ID", "share_id": None},
        }

        status = derive_purchasing_status(result)
        summary = build_purchasing_summary(result)

        # A fresh, untouched run with zero required purchasing lines and
        # no explicit regeneration yet is "Not Yet Generated" under the
        # new model -- the injected legacy "current" MyList status must
        # not flip this to some MyList-derived Ready/Failed reading.
        self.assertEqual(status, "Not Yet Generated")
        self.assertEqual(summary["status"], "Not Yet Generated")

        # The summary's digikey shape is built entirely from the new
        # model -- no trace of the legacy purchasing dict's fields leaks
        # into it.
        self.assertNotIn("list_id", summary["digikey"])
        self.assertNotIn("share_id", summary["digikey"])
        self.assertNotIn("created", summary["digikey"])
        self.assertEqual(summary["digikey"]["cart_link_url"], "")

    def test_legacy_mylist_shaped_response_survives_regeneration(self):
        result = self.run_workflow(csv_text=CLEAN_BOM_CSV)
        active_bom_state = result["_active_bom_state"]
        row_id = active_bom_state.rows()[0]["row_id"]
        self.make_digikey_sourced(active_bom_state, row_id)

        result["artifact_status"]["digikey_mylist"] = "current"
        result["purchasing"] = {
            "mouser": None,
            "digikey": {"created": True, "list_id": "OLD-LIST-ID", "share_id": None},
        }

        regenerate_purchasing(result)

        # A real regeneration against the new model succeeds normally;
        # the stale legacy keys are simply along for the ride, ignored.
        self.assertEqual(result["artifact_status"]["digikey_list_csv"], "current")
        self.assertEqual(result["artifact_status"]["digikey_cart_link"], "current")
        self.assertEqual(derive_purchasing_status(result), "Ready for Handoff")


if __name__ == "__main__":
    unittest.main()
